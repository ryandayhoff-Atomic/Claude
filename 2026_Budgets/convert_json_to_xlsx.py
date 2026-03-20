#!/usr/bin/env python3
"""Convert Google Sheets API JSON response to xlsx files using openpyxl."""
import json
import sys
import os
from openpyxl import Workbook


def json_to_xlsx(json_file_path, output_xlsx_path):
    """Convert a saved Sheets API JSON response to an xlsx file."""
    with open(json_file_path, 'r') as f:
        raw = json.load(f)

    # Handle the wrapper format: [{type: "text", text: "..."}]
    if isinstance(raw, list) and len(raw) > 0 and 'text' in raw[0]:
        data = json.loads(raw[0]['text'])
    else:
        data = raw

    # Navigate to sheets data
    sheets_data = data.get('outputs', data).get('tool_output', data).get('body', data).get('sheets', [])

    if not sheets_data:
        print(f"No sheets found in {json_file_path}")
        return False

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet_info in sheets_data:
        title = sheet_info.get('properties', {}).get('title', 'Sheet')
        # Excel sheet names max 31 chars
        title = title[:31]
        ws = wb.create_sheet(title=title)

        data_blocks = sheet_info.get('data', [])
        for block in data_blocks:
            row_data_list = block.get('rowData', [])
            for row_idx, row_data in enumerate(row_data_list, start=1):
                values = row_data.get('values', [])
                for col_idx, cell_data in enumerate(values, start=1):
                    formatted = cell_data.get('formattedValue', '')
                    if formatted:
                        cell = ws.cell(row=row_idx, column=col_idx)
                        # Try to convert numbers
                        try:
                            # Remove commas and try float
                            clean = formatted.replace(',', '').replace('$', '').replace('%', '')
                            if formatted.startswith('$') or formatted.endswith('%') or ',' in formatted:
                                val = float(clean)
                                cell.value = val
                                if formatted.startswith('$'):
                                    cell.number_format = '$#,##0.00'
                                elif formatted.endswith('%'):
                                    cell.value = val / 100
                                    cell.number_format = '0.00%'
                                else:
                                    if '.' in formatted:
                                        cell.number_format = '#,##0.00'
                                    else:
                                        cell.number_format = '#,##0'
                            else:
                                # Try int/float
                                if '.' in formatted:
                                    cell.value = float(formatted)
                                else:
                                    cell.value = int(formatted)
                        except (ValueError, TypeError):
                            cell.value = formatted

    wb.save(output_xlsx_path)
    print(f"Saved: {output_xlsx_path} ({len(sheets_data)} sheets)")
    return True


if __name__ == '__main__':
    if len(sys.argv) >= 3:
        json_to_xlsx(sys.argv[1], sys.argv[2])
    else:
        print("Usage: python convert_json_to_xlsx.py <input.json> <output.xlsx>")
